"""python-docx 版式辅助：中文字体、段落、标题、三线表、图与图题。"""
from docx import Document

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_CN, BODY_EN, HEAD_CN = "宋体", "Times New Roman", "黑体"


def set_font(run, cn=BODY_CN, en=BODY_EN, size=10.5, bold=False, color=None, italic=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), cn)
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(2.3), Cm(2.3)
    st = doc.styles["Normal"]
    st.font.name = BODY_EN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CN)
    pf = st.paragraph_format
    pf.line_spacing = 1.22
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    add_page_number(sec)
    return doc


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    set_font(run, size=9)


def para(doc, text, size=10.5, indent=True, align="justify", bold_prefix=None, space_after=2, color=None):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix:
        set_font(p.add_run(bold_prefix), cn=HEAD_CN, size=size, bold=True)
    set_font(p.add_run(text), size=size, color=color)
    return p


def heading(doc, text, level=1):
    size = {0: 17, 1: 13, 2: 11}[level]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 0 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt({0: 0, 1: 7, 2: 4}[level])
    p.paragraph_format.space_after = Pt({0: 3, 1: 3, 2: 2}[level])
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), cn=HEAD_CN, size=size, bold=True)
    return p


def _cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for edge, val in kw.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val.get("val", "single"))
        el.set(qn("w:sz"), str(val.get("sz", 8)))
        el.set(qn("w:color"), val.get("color", "000000"))
        b.append(el)


def table(doc, df, caption=None, col_widths=None, size=8.5, note=None):
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(3)
        cp.paragraph_format.space_after = Pt(1)
        cp.paragraph_format.keep_with_next = True
        set_font(cp.add_run(caption), cn=HEAD_CN, size=9, bold=True)
    t = doc.add_table(rows=len(df) + 1, cols=len(df.columns))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, c in enumerate(df.columns):
        cell = t.cell(0, j)
        cell.text = ""
        set_font(cell.paragraphs[0].add_run(str(c)), cn=HEAD_CN, size=size, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            set_font(cell.paragraphs[0].add_run(str(v)), size=size)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
    n = len(df)
    for j in range(len(df.columns)):
        _cell_border(t.cell(0, j), top={"sz": 12}, bottom={"sz": 6})
        _cell_border(t.cell(n, j), bottom={"sz": 12})
    for ri, row in enumerate(t.rows):
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.keep_with_next = ri < len(t.rows) - 1      # 整表不跨页
    if col_widths:
        for j, w in enumerate(col_widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    if note:
        np_ = doc.add_paragraph()
        np_.paragraph_format.space_after = Pt(3)
        set_font(np_.add_run(note), size=7.5, color="52606D")
    return t


def figure(doc, path, caption, width_cm=16.0, note=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(1 if note else 4)
    set_font(cp.add_run(caption), cn=HEAD_CN, size=9, bold=True)
    if note:
        np_ = doc.add_paragraph()
        np_.paragraph_format.space_after = Pt(4)
        set_font(np_.add_run(note), size=7.5, color="52606D")


def page_break(doc):
    doc.add_page_break()
